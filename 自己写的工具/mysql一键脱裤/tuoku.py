#@author:九世
#@file:tuoku.py
#@time:2019/10/7

import pymysql
import sys
import openpyxl
import warnings
import os

from colorama import Fore,Back,Style,init

warnings.filterwarnings("ignore")
init(wrap=True)
myxl=openpyxl.Workbook()
class Tuo(object):
    def __init__(self):
        self.host=sys.argv[1]
        self.port=sys.argv[2]
        self.username=sys.argv[3]
        self.password=sys.argv[4]
        self.database=sys.argv[5]

    def connect(self):
        print(Fore.YELLOW+'hosts:{} port:{} username:{} password:{} database:{}'.format(self.host,self.port,self.username,self.password,self.database))
        try:
            conn=pymysql.connect(host=self.host,port=int(self.port),user=self.username,password=self.password,database=self.database,charset='utf8')
            print(Fore.GREEN+'[+] 连接数据库成功')
            self.show_database(conn)
        except Exception as r:
            print(Fore.BLUE+'{}'.format(r))

    def show_database(self,conn):
        calc=0
        calc2=2
        bname=65
        print(Fore.YELLOW+'[+] 拥有以下表:')
        cursor=conn.cursor()
        cursor.execute("select table_name from information_schema.tables where table_schema='{}';".format(self.database))
        for u in cursor.fetchall():
            print(Fore.YELLOW+'表名:{}'.format(str(u[0])))
            myxl.create_sheet(index=calc,title='{}'.format(u[0]))
            msheet = myxl.get_sheet_by_name('{}'.format(u[0]))
            print('='*120)
            column_name=cursor.execute("select column_name from information_schema.columns where table_name='{}'".format(str(u[0])))
            for c in cursor.fetchall():
                zhi='{}1'.format(chr(bname))
                msheet[str(zhi)]='{}'.format(c[0])
                print(Fore.YELLOW+'字段名:{}'.format(c[0]))
                bname+=1
            print('=' * 120)
            bname=65

            cursor.execute("select * from {}".format(u[0]))

            jg=chr(bname+column_name)
            for s in cursor.fetchall():
                try:
                    for r in range(int(len(s))):
                        zhi='{}{}'.format(chr(bname),calc2)
                        msheet[zhi]=s[r]
                        bname+=1
                        if chr(bname)==jg:
                            bname=65
                            calc2+=1
            
                except Exception as r:
                    print(r)
                    pass

            calc+=1
        myxl.save('{}.xlsx'.format(self.database))
        print(Fore.GREEN+'[+] {}.xlsx保存成功'.format(self.database))

if __name__ == '__main__':
    try:
        obj=Tuo()
        obj.connect()
    except Exception as r:
        if 'list index out' in str(r):
            print(Fore.BLUE+'[!] run python tuoku.py [host] [port] [username] [password] [database]')
        else:
            print(Fore.BLUE+'{}'.format(r))